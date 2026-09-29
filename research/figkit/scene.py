"""One scene graph, two renderers: SVG for the paper, PPTX for editing.

THE POINT OF DOING IT THIS WAY

The obvious route to an editable figure is to draw an SVG and drop it
into PowerPoint. That gives you a picture, and "Convert to Shape" turns
it into several hundred ungrouped paths — every rounded rectangle becomes
an outline, every text run becomes curves. It is editable the way a
photograph of a whiteboard is editable.

So the figure is described once, as shapes, and rendered twice. The PPTX
export emits *native* PowerPoint objects — real rounded rectangles with
fill and line properties, real text frames you can retype, real
connectors you can drag. Only icons and screenshots arrive as images,
because they genuinely are images.

Coordinates are in points, origin top-left, and are shared by both
renderers: 1 unit = 1 SVG user unit = 1/72 inch in PowerPoint. A figure
960 wide is 13.3 inches, which is a 16:9 slide.
"""

from __future__ import annotations

import html
from dataclasses import dataclass, field
from pathlib import Path

from . import palette as P


# --------------------------------------------------------------------------
# nodes
# --------------------------------------------------------------------------

@dataclass
class Node:
    x: float
    y: float


@dataclass
class Box(Node):
    """A rounded container. `dashed` marks a grouping border, matching the
    references' convention that solid = a thing, dashed = a set of things."""
    w: float
    h: float
    fill: str = P.WHITE
    stroke: str = P.STROKE
    radius: float = 10
    dashed: bool = False
    width: float = 1.0
    shadow: bool = False


@dataclass
class Text(Node):
    text: str
    size: float = 11
    color: str = P.INK
    weight: str = "normal"      # normal | bold
    anchor: str = "start"       # start | middle | end
    mono: bool = False
    italic: bool = False
    w: float = 0                # >0 wraps into a PPT text frame of this width
    line: float = 1.35


@dataclass
class Icon(Node):
    spec: str
    size: float = 20
    color: str = "#334155"


@dataclass
class Image(Node):
    path: str
    w: float
    h: float
    radius: float = 6


@dataclass
class Arrow(Node):
    x2: float
    y2: float
    color: str = P.INK_FAINT
    width: float = 1.6
    dashed: bool = False
    head: bool = True


@dataclass
class Path(Node):
    """An orthogonal route through a list of points.

    Straight arrows can only say "A then B". A control point is defined by
    what comes *back* from it, so any figure of a mediation layer needs
    return paths, and those have to turn corners. `x`/`y` are the first
    point; `pts` carries the rest.
    """
    pts: list[tuple[float, float]] = field(default_factory=list)
    color: str = P.INK_FAINT
    width: float = 1.6
    dashed: bool = False
    head: bool = True
    radius: float = 6


@dataclass
class Scene:
    w: float
    h: float
    bg: str = P.WHITE
    nodes: list[Node] = field(default_factory=list)

    def add(self, *ns: Node) -> "Scene":
        self.nodes.extend(ns)
        return self

    # -- composite helpers -------------------------------------------------
    def panel(self, x, y, w, h, title=None, tint="blue", dashed=False,
              title_size=11.5):
        """A tinted stage container with an optional title in the tint's
        own darker text colour — the references' main structural unit."""
        self.add(Box(x, y, w, h, fill=P.TINT[tint], stroke=P.TINT_EDGE[tint],
                     radius=12, dashed=dashed))
        if title:
            self.add(Text(x + 12, y + 17, title, size=title_size, weight="bold",
                          color=P.TINT_TEXT[tint]))
        return self

    def card(self, x, y, w, h, label=None, size=9.5, icon=None, tint=None):
        """A white card on a tinted panel. Icon sits left, label beside it."""
        self.add(Box(x, y, w, h, fill=P.WHITE if tint is None else P.TINT[tint],
                     stroke=P.STROKE_SOFT, radius=7))
        tx = x + 9
        if icon:
            self.add(Icon(x + 9, y + h / 2 - 7, icon, size=14))
            tx = x + 29
        if label:
            self.add(Text(tx, y + h / 2 + size * 0.36, label, size=size))
        return self

    def chip(self, x, y, w, h, text, fill=P.ACCENT, fg=P.WHITE, size=9.5,
             mono=False, radius=None):
        """A saturated callout. Reserved for content that carries meaning:
        the task statement, the emitted action, the verdict."""
        self.add(Box(x, y, w, h, fill=fill, stroke=fill,
                     radius=h / 2 if radius is None else radius))
        self.add(Text(x + w / 2, y + h / 2 + size * 0.36, text, size=size,
                      color=fg, anchor="middle", mono=mono))
        return self

    def code(self, x, y, w, h, text, size=9):
        self.add(Box(x, y, w, h, fill=P.MONO_BG, stroke=P.STROKE_SOFT, radius=5))
        self.add(Text(x + 8, y + h / 2 + size * 0.36, text, size=size,
                      color=P.MONO_FG, mono=True))
        return self

    def verdict(self, x, y, ok, label=None, size=9.5):
        """Always a mark AND a word: reviewers print in greyscale."""
        col = P.PASS if ok else P.FAIL
        self.add(Icon(x, y - 6, "tabler:circle-check" if ok else "tabler:circle-x",
                      size=13, color=col))
        if label:
            self.add(Text(x + 17, y + size * 0.36 - 1, label, size=size, color=col))
        return self


# --------------------------------------------------------------------------
# SVG renderer  (paper)
# --------------------------------------------------------------------------

def _esc(s: str) -> str:
    return html.escape(str(s), quote=False)


def _rounded_path(pts: list[tuple[float, float]], r: float) -> str:
    """Polyline with the corners filleted.

    Square corners on a routing line read as a table border rather than a
    route; a small radius is what makes a reader follow it. The radius is
    clamped to half the shorter adjacent segment so a tight corner
    degenerates gracefully instead of overshooting.
    """
    if len(pts) < 2:
        return ""
    d = [f"M {pts[0][0]},{pts[0][1]}"]
    for i in range(1, len(pts) - 1):
        (x0, y0), (x1, y1), (x2, y2) = pts[i - 1], pts[i], pts[i + 1]
        l1 = max(1e-6, ((x1 - x0) ** 2 + (y1 - y0) ** 2) ** 0.5)
        l2 = max(1e-6, ((x2 - x1) ** 2 + (y2 - y1) ** 2) ** 0.5)
        rr = min(r, l1 / 2, l2 / 2)
        ax, ay = x1 - (x1 - x0) / l1 * rr, y1 - (y1 - y0) / l1 * rr
        bx, by = x1 + (x2 - x1) / l2 * rr, y1 + (y2 - y1) / l2 * rr
        d.append(f"L {ax},{ay}")
        d.append(f"Q {x1},{y1} {bx},{by}")
    d.append(f"L {pts[-1][0]},{pts[-1][1]}")
    return " ".join(d)


def to_svg(sc: Scene) -> str:
    from .icons import inline

    out = [
        f'<svg xmlns="http://www.w3.org/2000/svg" '
        f'xmlns:xlink="http://www.w3.org/1999/xlink" '
        f'width="{sc.w}" height="{sc.h}" viewBox="0 0 {sc.w} {sc.h}" '
        f'font-family="{P.SANS}">',
        '<defs><marker id="ah" viewBox="0 0 10 10" refX="8.5" refY="5" '
        'markerWidth="6" markerHeight="6" orient="auto-start-reverse">'
        '<path d="M0,1 L9,5 L0,9 z" fill="context-stroke"/></marker>'
        '<filter id="sh" x="-20%" y="-20%" width="140%" height="140%">'
        '<feDropShadow dx="0" dy="1.5" stdDeviation="2.5" '
        'flood-color="#0F172A" flood-opacity="0.07"/></filter></defs>',
        f'<rect width="{sc.w}" height="{sc.h}" fill="{sc.bg}"/>',
    ]

    for n in sc.nodes:
        if isinstance(n, Box):
            dash = ' stroke-dasharray="5 4"' if n.dashed else ""
            filt = ' filter="url(#sh)"' if n.shadow else ""
            out.append(
                f'<rect x="{n.x}" y="{n.y}" width="{n.w}" height="{n.h}" '
                f'rx="{n.radius}" fill="{n.fill}" stroke="{n.stroke}" '
                f'stroke-width="{n.width}"{dash}{filt}/>')

        elif isinstance(n, Text):
            fam = P.MONO if n.mono else P.SANS
            style = ' font-style="italic"' if n.italic else ""
            out.append(
                f'<text x="{n.x}" y="{n.y}" font-size="{n.size}" '
                f'fill="{n.color}" font-weight="{n.weight}" '
                f'text-anchor="{n.anchor}" font-family="{fam}"{style}>'
                f'{_esc(n.text)}</text>')

        elif isinstance(n, Icon):
            out.append(
                f'<g transform="translate({n.x},{n.y})">'
                f'<svg width="{n.size}" height="{n.size}" overflow="visible">'
                f'{inline(n.spec, n.color)}</svg></g>')

        elif isinstance(n, Image):
            # Base64 rather than a file:// href: rsvg-convert and most
            # PDF toolchains refuse to load external references, and they
            # do it silently — you get a correctly-sized blank rectangle.
            # Inlining also makes the .svg self-contained, which is what
            # you want handing a figure to a co-author.
            import base64
            import mimetypes
            p = Path(n.path)
            mime = mimetypes.guess_type(p.name)[0] or "image/png"
            b64 = base64.b64encode(p.read_bytes()).decode()
            out.append(
                f'<image x="{n.x}" y="{n.y}" width="{n.w}" height="{n.h}" '
                f'preserveAspectRatio="xMidYMid slice" '
                f'xlink:href="data:{mime};base64,{b64}"/>')

        elif isinstance(n, Arrow):
            dash = ' stroke-dasharray="4 3"' if n.dashed else ""
            head = ' marker-end="url(#ah)"' if n.head else ""
            out.append(
                f'<line x1="{n.x}" y1="{n.y}" x2="{n.x2}" y2="{n.y2}" '
                f'stroke="{n.color}" stroke-width="{n.width}" '
                f'stroke-linecap="round"{dash}{head}/>')

        elif isinstance(n, Path):
            dash = ' stroke-dasharray="4 3"' if n.dashed else ""
            head = ' marker-end="url(#ah)"' if n.head else ""
            pts = [(n.x, n.y)] + list(n.pts)
            d = _rounded_path(pts, n.radius)
            out.append(
                f'<path d="{d}" fill="none" stroke="{n.color}" '
                f'stroke-width="{n.width}" stroke-linecap="round" '
                f'stroke-linejoin="round"{dash}{head}/>')

    out.append("</svg>")
    return "\n".join(out)


# --------------------------------------------------------------------------
# PPTX renderer  (hand-editing)
# --------------------------------------------------------------------------

def to_pptx(sc: Scene, dest: Path) -> Path:
    """Native PowerPoint shapes, not a pasted picture.

    Rounded rectangles keep their corner radius as an adjustment handle,
    text stays text, arrows stay connectors. Icons and screenshots are
    placed as PNGs because that is what they are.
    """
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Pt

    from .icons import png as icon_png

    def emu(v: float) -> int:
        return Emu(int(round(v * 12700)))       # 1 pt = 12700 EMU

    def rgb(c: str) -> RGBColor:
        return RGBColor.from_string(c.lstrip("#").upper()[:6])

    prs = Presentation()
    prs.slide_width, prs.slide_height = emu(sc.w), emu(sc.h)
    slide = prs.slides.add_slide(prs.slide_layouts[6])   # blank

    if sc.bg and sc.bg.upper() != "#FFFFFF":
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                    emu(sc.w), emu(sc.h))
        bg.fill.solid(); bg.fill.fore_color.rgb = rgb(sc.bg)
        bg.line.fill.background()

    for n in sc.nodes:
        if isinstance(n, Box):
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE if n.radius > 0 else MSO_SHAPE.RECTANGLE,
                emu(n.x), emu(n.y), emu(n.w), emu(n.h))
            # PowerPoint stores corner radius as a fraction of the short
            # side, so a fixed pt radius has to be converted per shape.
            if n.radius > 0:
                try:
                    shape.adjustments[0] = min(0.5, n.radius / max(1.0, min(n.w, n.h)))
                except (IndexError, ValueError):
                    pass
            shape.fill.solid(); shape.fill.fore_color.rgb = rgb(n.fill)
            shape.line.color.rgb = rgb(n.stroke)
            shape.line.width = Pt(n.width)
            if n.dashed:
                from pptx.enum.dml import MSO_LINE_DASH_STYLE
                shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            shape.shadow.inherit = False
            if shape.has_text_frame:
                shape.text_frame.text = ""

        elif isinstance(n, Text):
            w = n.w if n.w else max(60.0, len(n.text) * n.size * 0.62)
            x = n.x if n.anchor == "start" else (
                n.x - w / 2 if n.anchor == "middle" else n.x - w)
            tb = slide.shapes.add_textbox(emu(x), emu(n.y - n.size * 1.15),
                                          emu(w), emu(n.size * n.line * 1.6))
            tf = tb.text_frame
            tf.word_wrap = bool(n.w)
            tf.margin_left = tf.margin_right = 0
            tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.TOP
            p = tf.paragraphs[0]
            p.alignment = {"start": PP_ALIGN.LEFT, "middle": PP_ALIGN.CENTER,
                           "end": PP_ALIGN.RIGHT}[n.anchor]
            r = p.add_run(); r.text = n.text
            f = r.font
            f.size = Pt(n.size); f.bold = (n.weight == "bold")
            f.italic = n.italic
            f.name = P.MONO_PPT if n.mono else P.SANS_PPT
            f.color.rgb = rgb(n.color)

        elif isinstance(n, Icon):
            slide.shapes.add_picture(str(icon_png(n.spec, n.color, 256)),
                                     emu(n.x), emu(n.y),
                                     emu(n.size), emu(n.size))

        elif isinstance(n, Image):
            if Path(n.path).exists():
                slide.shapes.add_picture(n.path, emu(n.x), emu(n.y),
                                         emu(n.w), emu(n.h))

        elif isinstance(n, Arrow):
            from pptx.enum.shapes import MSO_CONNECTOR
            cx = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, emu(n.x), emu(n.y), emu(n.x2), emu(n.y2))
            cx.line.color.rgb = rgb(n.color)
            cx.line.width = Pt(n.width)

        elif isinstance(n, Path):
            # One connector per segment rather than a freeform: the segments
            # stay individually draggable, which is the point of the
            # editable export.
            from pptx.enum.shapes import MSO_CONNECTOR
            pts = [(n.x, n.y)] + list(n.pts)
            for (ax, ay), (bx, by) in zip(pts, pts[1:]):
                cx = slide.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT, emu(ax), emu(ay), emu(bx), emu(by))
                cx.line.color.rgb = rgb(n.color)
                cx.line.width = Pt(n.width)

    dest.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dest))
    return dest


# --------------------------------------------------------------------------
# export
# --------------------------------------------------------------------------

def export(sc: Scene, stem: Path, *, pdf: bool = True, pptx: bool = True,
           png: bool = False) -> dict[str, Path]:
    """Write .svg (master), .pdf (LaTeX), .pptx (hand-editing)."""
    import subprocess

    stem.parent.mkdir(parents=True, exist_ok=True)
    made: dict[str, Path] = {}

    svg_path = stem.with_suffix(".svg")
    svg_path.write_text(to_svg(sc))
    made["svg"] = svg_path

    def convert(args: list[str]) -> bool:
        try:
            subprocess.run(args, check=True, capture_output=True)
            return True
        except (FileNotFoundError, subprocess.CalledProcessError):
            return False

    if pdf:
        p = stem.with_suffix(".pdf")
        if convert(["rsvg-convert", "-f", "pdf", "-o", str(p), str(svg_path)]) \
           or convert(["cairosvg", str(svg_path), "-o", str(p)]):
            made["pdf"] = p
    if png:
        p = stem.with_suffix(".png")
        if convert(["rsvg-convert", "-z", "2", "-o", str(p), str(svg_path)]) \
           or convert(["cairosvg", str(svg_path), "-o", str(p), "-s", "2"]):
            made["png"] = p
    if pptx:
        made["pptx"] = to_pptx(sc, stem.with_suffix(".pptx"))
    return made
